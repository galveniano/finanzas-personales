"""Cobros, libro, resumen anual, clientes, nóminas y supuestos del flujo de ingresos (datos inventados, años raros)."""
import io
from datetime import date, timedelta
from decimal import Decimal as D

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from finanzas import actividad, api, db, prevision
from finanzas.importers.nomina_pdf import LecturaNomina
from finanzas.main import app
from finanzas.models import Categoria, Cuenta, Factura, GastoAutonomo, Movimiento, Nomina
from test_nomina_pdf import RECIBO, _pdf


def _factura(c, numero, cliente, fecha, base, iva=21, ret=15, **extra):
    r = c.post("/api/autonomo/facturas", json={"numero": numero, "cliente": cliente, "fecha": fecha, "base": base,
                                               "tipo_iva": iva, "tipo_retencion": ret, **extra})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _cuenta_con(movimientos: list[tuple[str, float, str | None, str]], nombre="Cuenta cobros ejemplo") -> int:
    """Cuenta con movimientos (fecha, importe, categoría, concepto); devuelve su id."""
    with db.SessionLocal() as s:
        categorias = {c.nombre: c.id for c in s.scalars(select(Categoria))}
        cuenta = Cuenta(nombre=nombre)
        s.add(cuenta)
        s.flush()
        for i, (fecha, importe, categoria, concepto) in enumerate(movimientos):
            s.add(Movimiento(cuenta_id=cuenta.id, fecha=date.fromisoformat(fecha), importe=D(str(importe)),
                             categoria_id=categorias.get(categoria) if categoria else None, concepto=concepto,
                             huella=f"{nombre}-{i}"))
        s.commit()
        return cuenta.id


def test_sugerencias_de_cobro_y_conciliar():
    with TestClient(app) as c:
        f1 = _factura(c, "COB-1", "Cobros Test SL", "2012-03-01", 1000)  # total 1.060
        f2 = _factura(c, "COB-2", "Cobros Test SL", "2012-03-05", 1000)  # también 1.060
        f3 = _factura(c, "COB-3", "Cobros Test SL", "2012-04-01", 500, iva=0, ret=0)  # total 500
        cuenta = _cuenta_con([
            ("2012-03-10", 1060.60, None, "TRANSFERENCIA RECIBIDA"),  # cuadra (±1 €) pero sin categoría
            ("2012-03-12", 1060.00, "Cobro de facturas", "COBRO FACTURA"),  # se prefiere: es un cobro de facturas
            ("2012-02-01", 500.00, None, "ANTES DE LA FACTURA"),  # anterior a la factura 3: no vale
            ("2012-04-15", 503.00, None, "NO CUADRA"),  # fuera del margen de 1 €
        ])
        d = c.get("/api/autonomo?anio=2012").json()
        assert d["por_cobrar"] == {"total": 2620.0, "facturas": 3, "mas_antigua_dias": (date.today() - date(2012, 3, 1)).days}
        por_factura = {x["id"]: x for x in d["facturas"]}
        assert not por_factura[f1]["cobrada"] and por_factura[f1]["dias_pendiente"] == d["por_cobrar"]["mas_antigua_dias"]
        sug = {s_["factura_id"]: s_ for s_ in d["sugerencias_cobro"]}
        assert sug[f1]["fecha"] == "2012-03-12" and sug[f1]["importe"] == 1060
        assert sug[f2]["fecha"] == "2012-03-10" and sug[f2]["importe"] == 1060.6  # el otro ingreso ya está usado
        assert f3 not in sug and len(sug) == 2

        # Los avisos de Inicio: facturas de hace más de 45 días sin cobrar
        with db.SessionLocal() as s:
            avisos = actividad.avisos(s)
        assert avisos and all(a["ir"] == "/ingresos" and "sin cobrar" in a["texto"] for a in avisos)

        assert c.post("/api/autonomo/cobros/conciliar").json()["marcadas"] == 2
        d = c.get("/api/autonomo?anio=2012").json()
        por_factura = {x["id"]: x for x in d["facturas"]}
        assert por_factura[f1]["fecha_cobro"] == "2012-03-12" and por_factura[f1]["cobrada"]
        assert por_factura[f1]["dias_pendiente"] is None and por_factura[f2]["fecha_cobro"] == "2012-03-10"
        assert d["por_cobrar"]["facturas"] == 1 and not d["sugerencias_cobro"]
        # Cobrada hoy (sin fecha) y con fecha
        assert c.post(f"/api/autonomo/facturas/{f3}/cobrada", json={}).json()["fecha_cobro"] == date.today().isoformat()
        assert c.post(f"/api/autonomo/facturas/{f3}/cobrada", json={"fecha": "2012-05-02"}).json()["fecha_cobro"] == "2012-05-02"
        assert c.get("/api/autonomo?anio=2012").json()["por_cobrar"]["facturas"] == 0
        assert c.post("/api/autonomo/facturas/999999/cobrada", json={}).status_code == 404
        for fid in (f1, f2, f3):
            c.delete(f"/api/autonomo/facturas/{fid}")
        c.delete(f"/api/cuentas/{cuenta}")


def test_avisos_de_numeracion():
    with TestClient(app) as c:
        ids = [_factura(c, n, "Numeración Test", f, 100) for n, f in
               (("2015-001", "2015-02-01"), ("2015-002", "2015-03-01"), ("2015-004", "2015-04-01"))]
        avisos = c.get("/api/autonomo?anio=2015").json()["avisos_numeracion"]
        assert len(avisos) == 1 and "2015-003" in avisos[0] and "2015-004" in avisos[0]
        # Fecha fuera de orden respecto al número
        ids.append(_factura(c, "2015-003", "Numeración Test", "2015-05-01", 100))
        avisos = c.get("/api/autonomo?anio=2015").json()["avisos_numeracion"]
        assert len(avisos) == 2 and "2015-005" in avisos[1]
        for fid in ids:
            c.delete(f"/api/autonomo/facturas/{fid}")
        assert c.get("/api/autonomo?anio=2015").json()["avisos_numeracion"] == []


def test_libro_xlsx_para_la_gestoria():
    from openpyxl import load_workbook
    with TestClient(app) as c:
        f1 = _factura(c, "LIB-1", "Libro Test SL", "2012-01-15", 1000, fecha_cobro="2012-02-01")
        f2 = _factura(c, "LIB-2", "Libro Test SL", "2012-05-15", 2000)
        g = c.post("/api/autonomo/gastos", json={"fecha": "2012-02-10", "proveedor": "Gestoría Test", "concepto": "Trimestre",
                                                 "categoria": "gestoria", "base": 100, "tipo_iva": 21, "deducible_pct": 50})
        assert g.status_code == 200
        r = c.get("/api/autonomo/libro.xlsx?anio=2012")
        assert r.status_code == 200 and "libro-facturas-2012.xlsx" in r.headers["content-disposition"]
        libro = load_workbook(io.BytesIO(r.content))
        assert libro.sheetnames == ["Emitidas", "Gastos"]
        emitidas = [list(f) for f in libro["Emitidas"].iter_rows(values_only=True)]
        assert emitidas[0][:6] == ["Número", "Fecha", "Cliente", "NIF", "Concepto", "Base"]
        assert emitidas[1][0] == "LIB-1" and emitidas[1][5] == 1000 and emitidas[1][7] == 210 and emitidas[1][10] == 1060
        assert emitidas[1][11].date() == date(2012, 2, 1) if hasattr(emitidas[1][11], "date") else emitidas[1][11] == date(2012, 2, 1)
        totales = {f[0]: f for f in emitidas if f[0] and str(f[0]).startswith("Total")}
        assert totales["Total 1T"][5] == 1000 and totales["Total 2T"][5] == 2000 and totales["Total 2012"][5] == 3000
        assert "Total 3T" not in totales  # sin facturas en ese trimestre no hay fila
        gastos = [list(f) for f in libro["Gastos"].iter_rows(values_only=True)]
        assert gastos[0][-2:] == ["Base deducible", "IVA deducible"]
        assert gastos[1][1] == "Gestoría Test" and gastos[1][4] == 100 and gastos[1][8] == 50 and gastos[1][9] == 10.5
        gid = next(x["id"] for x in c.get("/api/autonomo?anio=2012").json()["gastos"] if x["proveedor"] == "Gestoría Test")
        for fid in (f1, f2):
            c.delete(f"/api/autonomo/facturas/{fid}")
        c.delete(f"/api/autonomo/gastos/{gid}")


def test_resumen_anual_390_y_347():
    with TestClient(app) as c:
        ids = [_factura(c, "AN-1", "Grande Test SL", "2012-02-01", 3000),  # 3.630 con IVA: pasa de 3.005,06
               _factura(c, "AN-2", "Pequeño Test", "2012-02-02", 1000, iva=10, ret=0),
               _factura(c, "AN-3", "Extranjero Test", "2012-03-01", 4000, iva=0, ret=0)]
        c.post("/api/autonomo/gastos", json={"fecha": "2012-03-10", "proveedor": "Proveedor Grande", "base": 3000,
                                             "tipo_iva": 21, "deducible_pct": 100})
        c.post("/api/autonomo/gastos", json={"fecha": "2012-03-11", "proveedor": "Luz casa", "base": 100,
                                             "tipo_iva": 21, "deducible_pct": 30})
        d = c.get("/api/autonomo/anual?anio=2012").json()
        por_tipo = {r["tipo"]: r for r in d["repercutido"]}
        assert [r["tipo"] for r in d["repercutido"]] == [21, 10, 4, 0]
        assert por_tipo[21] == {"tipo": 21, "base": 3000, "cuota": 630} and por_tipo[10]["cuota"] == 100
        assert por_tipo[0]["base"] == 4000 and por_tipo[4]["base"] == 0
        assert d["base_total"] == 8000 and d["iva_repercutido"] == 730 and d["retenciones"] == 450
        assert d["iva_soportado"] == 630 + 6.3 and d["base_soportada"] == 3030
        assert d["presentado_303"] is None and d["calculado_303"] == round(730 - 636.3, 2)
        terceros = {t["nombre"]: t for t in d["terceros_347"]}
        assert set(terceros) == {"Grande Test SL", "Extranjero Test", "Proveedor Grande"}
        assert terceros["Grande Test SL"]["importe"] == 3630 and terceros["Grande Test SL"]["con_retencion"]
        assert terceros["Extranjero Test"]["tipo"] == "cliente" and not terceros["Extranjero Test"]["con_retencion"]
        assert terceros["Proveedor Grande"] == {"nombre": "Proveedor Grande", "nif": "", "tipo": "proveedor", "importe": 3630,
                                                "operaciones": 1, "con_retencion": False}
        # Con un 303 presentado, el resumen lo compara con el calculado
        c.post("/api/declaraciones", json={"modelo": "303", "ejercicio": 2012, "periodo": "1T", "importe": 90})
        d = c.get("/api/autonomo/anual?anio=2012").json()
        assert d["presentado_303"] == 90 and d["trimestres"][0]["presentado"] == 90 and d["trimestres"][1]["presentado"] is None
        decl = next(x["id"] for x in c.get("/api/declaraciones").json()["declaraciones"] if x["ejercicio"] == 2012 and x["modelo"] == "303")
        c.delete(f"/api/declaraciones/{decl}")
        for fid in ids:
            c.delete(f"/api/autonomo/facturas/{fid}")
        for g in c.get("/api/autonomo?anio=2012").json()["gastos"]:
            c.delete(f"/api/autonomo/gastos/{g['id']}")


def test_renombrar_y_borrar_cliente():
    with TestClient(app) as c:
        sup = {"clientes": [{"nombre": "Cliente Viejo Test", "tarifa_hora": 30, "horas_dia": 8, "dias_mes": 20, "iva": 21, "retencion": 15}]}
        assert c.put("/api/prevision/supuestos", json=sup).status_code == 200
        c.put("/api/prevision/dias", json={"cliente": "Cliente Viejo Test", "mes": "2040-01", "dias": 5})
        fid = _factura(c, "REN-1", "Cliente Viejo Test", "2012-06-01", 100)
        assert c.patch("/api/facturacion/clientes/Cliente%20Viejo%20Test", json={"nuevo_nombre": " Cliente Nuevo Test "}).status_code == 200
        assert c.get("/api/autonomo?anio=2012").json()["facturas"][0]["cliente"] == "Cliente Nuevo Test"
        cfg = c.get("/api/prevision").json()["supuestos"]
        assert cfg["clientes"][0]["nombre"] == "Cliente Nuevo Test" and cfg["dias_planificados"] == {"Cliente Nuevo Test": {"2040-01": 5}}
        clientes = {x["nombre"]: x for x in c.get("/api/facturacion").json()["clientes"]}
        assert "Cliente Viejo Test" not in clientes and clientes["Cliente Nuevo Test"]["facturas"] == 1
        # Nombre repetido o vacío
        _factura(c, "REN-2", "Otro Test", "2012-06-02", 100)
        assert c.patch("/api/facturacion/clientes/Otro%20Test", json={"nuevo_nombre": "Cliente Nuevo Test"}).status_code == 400
        assert c.patch("/api/facturacion/clientes/Otro%20Test", json={"nuevo_nombre": "  "}).status_code == 400
        assert c.patch("/api/facturacion/clientes/No%20Existe", json={"nuevo_nombre": "X"}).status_code == 404
        # Con facturas no se borra; sin ellas sí, y desaparece también de los supuestos
        r = c.delete("/api/facturacion/clientes/Cliente%20Nuevo%20Test")
        assert r.status_code == 400 and "1 factura" in r.json()["detail"]
        c.delete(f"/api/autonomo/facturas/{fid}")
        assert c.delete("/api/facturacion/clientes/Cliente%20Nuevo%20Test").status_code == 200
        assert "Cliente Nuevo Test" not in {x["nombre"] for x in c.get("/api/facturacion").json()["clientes"]}
        cfg = c.get("/api/prevision").json()["supuestos"]
        assert cfg["clientes"] == [] and cfg["dias_planificados"] == {}
        otro = next(x["id"] for x in c.get("/api/autonomo?anio=2012").json()["facturas"])
        c.delete(f"/api/autonomo/facturas/{otro}")
        c.delete("/api/facturacion/clientes/Otro%20Test")


def test_textos_recortados_a_la_columna():
    with TestClient(app) as c:
        largo = "N" * 150
        fid = _factura(c, "X" * 60, largo, "2012-07-01", 10)
        x = next(f for f in c.get("/api/autonomo?anio=2012").json()["facturas"] if f["id"] == fid)
        assert len(x["numero"]) == 40 and len(x["cliente"]) == 120
        assert c.put(f"/api/facturacion/clientes/{largo[:120]}", json={"nif": "A" * 40}).status_code == 200
        assert len(next(k["nif"] for k in c.get("/api/facturacion").json()["clientes"] if k["nombre"] == largo[:120])) == 20
        r = c.post("/api/nominas", json={"empresa": "E" * 100, "fecha": "2012-07-31", "bruto": 10, "retencion_irpf": 1,
                                          "seguridad_social": 1, "neto": 8})
        assert r.status_code == 200
        nom = next(x for x in c.get("/api/nominas?anio=2012").json()["nominas"] if x["fecha"] == "2012-07-31")
        assert len(nom["empresa"]) == 80
        c.delete(f"/api/nominas/{nom['id']}")
        c.delete(f"/api/autonomo/facturas/{fid}")
        c.delete(f"/api/facturacion/clientes/{largo[:120]}")


def test_editar_gasto():
    with TestClient(app) as c:
        c.post("/api/autonomo/gastos", json={"fecha": "2012-09-01", "proveedor": "Editable", "base": 50, "tipo_iva": 21})
        gid = next(g["id"] for g in c.get("/api/autonomo?anio=2012").json()["gastos"] if g["proveedor"] == "Editable")
        r = c.put(f"/api/autonomo/gastos/{gid}", json={"fecha": "2012-09-02", "proveedor": "P" * 200, "concepto": "Nuevo",
                                                        "categoria": "software", "base": 60, "tipo_iva": 0, "deducible_pct": 30})
        assert r.status_code == 200
        g = next(g for g in c.get("/api/autonomo?anio=2012").json()["gastos"] if g["id"] == gid)
        assert g["fecha"] == "2012-09-02" and len(g["proveedor"]) == 120 and g["categoria"] == "software"
        assert g["base"] == 60 and g["cuota_iva"] == 0 and g["deducible_pct"] == 30
        assert c.put("/api/autonomo/gastos/999999", json={"fecha": "2012-09-02", "base": 1}).status_code == 404
        c.delete(f"/api/autonomo/gastos/{gid}")


def test_cuota_de_autonomos_del_banco_como_gasto():
    with TestClient(app) as c:
        cuenta = _cuenta_con([("2014-01-31", -310, "Cuota autónomos", "TGSS COTIZACION 0114"),
                              ("2014-02-28", -310, None, "SEGURIDAD SOCIAL CUOTA"),
                              ("2014-02-28", -45, None, "CAFETERIA")], nombre="Cuenta TGSS ejemplo")
        # Enero ya está apuntado (con dos días de diferencia); febrero no
        c.post("/api/autonomo/gastos", json={"fecha": "2014-02-02", "proveedor": "TGSS", "categoria": "cuota_reta",
                                             "base": 310, "tipo_iva": 0})
        d = c.get("/api/autonomo/cuota-tgss?anio=2014").json()
        assert d == {"cargos": [{"fecha": "2014-02-28", "importe": 310}], "n": 1, "total": 310}
        r = c.post("/api/autonomo/cuota-tgss?anio=2014").json()
        assert r["ok"] and r["n"] == 1
        assert c.get("/api/autonomo/cuota-tgss?anio=2014").json()["n"] == 0
        nuevo = next(g for g in c.get("/api/autonomo?anio=2014").json()["gastos"] if g["fecha"] == "2014-02-28")
        assert nuevo["categoria"] == "cuota_reta" and nuevo["base"] == 310 and nuevo["tipo_iva"] == 0 and nuevo["deducible_pct"] == 100
        assert "febrero" in nuevo["concepto"]
        for g in c.get("/api/autonomo?anio=2014").json()["gastos"]:
            c.delete(f"/api/autonomo/gastos/{g['id']}")
        c.delete(f"/api/cuentas/{cuenta}")


def test_nomina_a_mano_y_pdf_del_mismo_mes_no_se_duplican(monkeypatch):
    with TestClient(app) as c:
        r = c.post("/api/nominas", json={"fecha": "2013-05-31", "bruto": 3000, "retencion_irpf": 500, "seguridad_social": 190, "neto": 2310})
        assert r.status_code == 200
        antes = c.get("/api/nominas?anio=2013").json()["nominas"]
        assert len(antes) == 1 and antes[0]["empresa"] == "Empresa"  # sin un nombre real por defecto
        # El PDF de ese mes (otra empresa, otras cifras) actualiza la fila en vez de añadir otra
        pdf = _pdf(RECIBO.replace("2026", "2013").replace("09/", "05/").replace("30/05", "31/05"))
        r = c.post("/api/nominas/pdf", files=[("ficheros", ("mayo.pdf", pdf, "application/pdf"))]).json()["resultados"][0]
        assert r["ok"] and r["mensaje"] == "Nómina de mayo de 2013 (actualizada)", r
        despues = c.get("/api/nominas?anio=2013").json()["nominas"]
        assert len(despues) == 1 and despues[0]["id"] == antes[0]["id"] and despues[0]["empresa"] == "Ejemplo Sistemas"
        assert despues[0]["neto"] == 2485.69 and despues[0]["tiene_pdf"]
        # Una paga extra del mismo mes sí es otra fila
        extra = LecturaNomina(empresa="Ejemplo", fecha=date(2013, 5, 31), bruto=D("1000"), seguridad_social=D("0"),
                              retencion_irpf=D("275"), neto=D("725"), paga_extra=True)
        monkeypatch.setattr(api, "_leer_nomina", lambda s, contenido, nombre: extra)
        r = c.post("/api/nominas/pdf", files=[("ficheros", ("extra.pdf", b"x", "application/pdf"))]).json()["resultados"][0]
        assert r["mensaje"] == "Paga extra de mayo de 2013"
        assert len(c.get("/api/nominas?anio=2013").json()["nominas"]) == 2
        for x in c.get("/api/nominas?anio=2013").json()["nominas"]:
            c.delete(f"/api/nominas/{x['id']}")


def test_guardar_supuestos_conserva_lo_de_la_renta():
    with TestClient(app) as c:
        with db.SessionLocal() as s:
            prevision.guardar(s, {**prevision.leer(s), "aportacion_pensiones_anio": 1500, "aportacion_ppes_anio": 1000,
                                  "fraccionar_renta": True, "rentas_ahorro_anio": 200, "imputacion_inmuebles_anio": 300,
                                  "dias_planificados": {"Cliente A": {"2040-02": 3}}})
        formulario = {"nomina": {"empresa": "E", "bruto_anual": 30000, "pagas": 14}, "clientes": [], "gastos_autonomo_mes": 50,
                      "gasto_habitual_mes": None, "meses_sin_facturar": [8]}
        assert c.put("/api/prevision/supuestos", json=formulario).status_code == 200
        with db.SessionLocal() as s:
            cfg = prevision.leer(s)
        assert cfg["aportacion_pensiones_anio"] == 1500 and cfg["aportacion_ppes_anio"] == 1000 and cfg["fraccionar_renta"] is True
        assert cfg["rentas_ahorro_anio"] == 200 and cfg["imputacion_inmuebles_anio"] == 300
        assert cfg["dias_planificados"] == {"Cliente A": {"2040-02": 3}} and cfg["nomina"]["bruto_anual"] == 30000
        assert cfg["gastos_autonomo_mes"] == 50 and cfg["meses_sin_facturar"] == [8]
        # nomina: null sí quita la nómina
        assert c.put("/api/prevision/supuestos", json={**formulario, "nomina": None}).status_code == 200
        with db.SessionLocal() as s:
            cfg = prevision.leer(s)
        assert cfg["nomina"] is None and cfg["fraccionar_renta"] is True


def test_calculo_de_nomina_con_la_retencion_que_quieres():
    with TestClient(app) as c:
        normal = c.get("/api/nominas/calculo?bruto_anual=30000").json()
        alto = c.get("/api/nominas/calculo?bruto_anual=30000&tipo_irpf=30").json()
        assert alto["tipo_irpf"] == 30 and alto["neto_mes"] < normal["neto_mes"] and alto["ss_anual"] == normal["ss_anual"]
        assert c.get("/api/nominas/calculo?bruto_anual=30000&tipo_irpf=60").status_code == 400


def test_nominas_por_anio_y_mes_a_mes_con_el_banco():
    anio = date.today().year - 1
    with TestClient(app) as c:
        r = c.post("/api/nominas", json={"empresa": "Mes a mes", "fecha": f"{anio}-03-31", "bruto": 3000, "retencion_irpf": 800,
                                          "seguridad_social": 200, "neto": 2000, "base_irpf": 2950, "especie": 40})
        assert r.status_code == 200
        cuenta = _cuenta_con([(f"{anio}-04-02", 2000, "Nómina", "NOMINA MARZO"),  # a primeros de mes: es la de marzo
                              (f"{anio}-05-28", 2500, "Nómina", "NOMINA MAYO")], nombre="Cuenta mes a mes")
        d = c.get(f"/api/nominas?anio={anio}").json()
        assert [x["fecha"] for x in d["nominas"]] == [f"{anio}-03-31"]  # solo las del año pedido
        assert d["nominas"][0]["base_irpf"] == 2950 and d["nominas"][0]["especie"] == 40
        assert len(d["meses"]) == 12 and d["meses"][0]["mes"] == f"{anio}-01"
        marzo, abril, mayo = d["meses"][2], d["meses"][3], d["meses"][4]
        assert marzo == {"mes": f"{anio}-03", "nomina_neto": 2000, "banco_importe": 2000, "banco_fecha": f"{anio}-04-02", "cuadra": True}
        assert abril["banco_importe"] is None and abril["nomina_neto"] is None and not abril["cuadra"]
        assert mayo["nomina_neto"] is None and mayo["banco_importe"] == 2500 and not mayo["cuadra"]  # falta la nómina de mayo
        assert c.get("/api/nominas?anio=2012").json()["meses"] == [m for m in c.get("/api/nominas?anio=2012").json()["meses"]]
        assert c.get(f"/api/nominas?anio={anio + 5}").json()["meses"] == []
        c.delete(f"/api/nominas/{d['nominas'][0]['id']}")
        c.delete(f"/api/cuentas/{cuenta}")


def test_retencion_recomendada(monkeypatch):
    hoy = date.today()
    if hoy.month == 12:
        pytest.skip("en diciembre no queda bruto por cobrar: no hay recomendación")
    with TestClient(app) as c:
        assert c.get("/api/nominas").json()["retencion_recomendada"] is None  # sin nómina ni supuestos
        c.put("/api/prevision/supuestos", json={"nomina": {"empresa": "E", "bruto_anual": 42000, "pagas": 14}})
        with db.SessionLocal() as s:
            x = Nomina(empresa="Recom", fecha=date(hoy.year, 1, 31), bruto=D("3000"), retencion_irpf=D("300"),
                       seguridad_social=D("190"), neto=D("2510"), tipo_irpf=D("10"))
            s.add(x)
            s.commit()
            nid = x.id
        try:
            # Con la previsión diciendo que la renta sale a pagar 1.000 €
            monkeypatch.setattr(prevision, "calcular", lambda s, meses=12: {"anios": [{"anio": hoy.year, "resultado": 1000.0}]})
            r = c.get("/api/nominas").json()["retencion_recomendada"]
            restante = sum(v for m, v in prevision._nomina_anual({"bruto_anual": 42000, "pagas": 14})["brutos"].items() if m > hoy.month)
            assert r["tipo_actual"] == 10 and r["resultado_previsto"] == 1000 and r["meses_restantes"] == 12 - hoy.month
            assert r["tipo_recomendado"] == round(10 + 1000 / restante * 100, 1) and 11 < r["tipo_recomendado"] <= 47
            # A devolver: no hay nada que pedir
            monkeypatch.setattr(prevision, "calcular", lambda s, meses=12: {"anios": [{"anio": hoy.year, "resultado": -400.0}]})
            assert c.get("/api/nominas").json()["retencion_recomendada"] is None
            # Otro año: tampoco
            assert c.get(f"/api/nominas?anio={hoy.year - 1}").json()["retencion_recomendada"] is None
        finally:
            c.delete(f"/api/nominas/{nid}")


def test_cobros_ignoran_traspasos_entre_cuentas():
    """Un traspaso entre dos cuentas tuyas del importe de la factura no es un cobro."""
    with TestClient(app) as c:
        fid = _factura(c, "TR-1", "Traspaso Test", "2012-11-01", 1000, iva=0, ret=0)
        a = _cuenta_con([("2012-11-05", -1000, None, "A AHORRO")], nombre="Cuenta traspaso A")
        b = _cuenta_con([("2012-11-05", 1000, None, "DE CORRIENTE")], nombre="Cuenta traspaso B")
        assert c.get("/api/autonomo?anio=2012").json()["sugerencias_cobro"] == []
        c.delete(f"/api/autonomo/facturas/{fid}")
        c.delete(f"/api/cuentas/{a}")
        c.delete(f"/api/cuentas/{b}")
        with db.SessionLocal() as s:
            assert s.scalar(select(Factura.id).where(Factura.numero == "TR-1")) is None
            assert s.scalar(select(GastoAutonomo.id).where(GastoAutonomo.fecha == date(2014, 2, 28))) is None
