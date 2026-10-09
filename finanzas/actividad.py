"""La actividad como autónomo más allá de registrar facturas: lo que te deben y su conciliación con el banco,
el libro de facturas para la gestoría, el resumen del año (390 y 347), la cuota de autónomos que carga el banco
como gasto, la numeración de las facturas y la edición de gastos."""
import io
import re
from datetime import date, timedelta
from decimal import Decimal

from fastapi import APIRouter, Depends, Response
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from finanzas import auth, db, declaraciones, facturacion, prevision
from finanzas.api import GastoAutonomoIn, _obtener, f, n
from finanzas.fechas import MESES
from finanzas.fiscal import autonomo
from finanzas.models import Factura, GastoAutonomo

router = APIRouter(prefix="/api", dependencies=[Depends(auth.requiere_sesion)])
SesionDB = Depends(db.get_session)
CERO = Decimal("0")
CENT = Decimal("0.01")
DIAS_AVISO_COBRO = 45  # a partir de aquí una factura sin cobrar es un aviso
UMBRAL_347 = Decimal("3005.06")  # operaciones con un tercero en el año (IVA incluido) que obligan al 347
TOLERANCIA_COBRO = 1.0  # euros de diferencia admitidos entre el ingreso del banco y el total de la factura
DIAS_CUOTA = 3  # días de margen entre el cargo del banco y la fecha del gasto apuntado


def _del_anio(s: Session, modelo, anio: int):
    return s.scalars(select(modelo).where(modelo.fecha >= date(anio, 1, 1), modelo.fecha < date(anio + 1, 1, 1))
                     .order_by(modelo.fecha, modelo.id)).all()


# --- Cobros pendientes --------------------------------------------------------

def estado_cobros(s: Session, facturas: list[Factura], hoy: date | None = None) -> dict:
    """Lo que te deben (facturas sin fecha de cobro, de cualquier año) y, para cada una, el ingreso del banco que
    parece su cobro: el primer ingreso en tus cuentas desde la fecha de la factura cuyo importe coincide con el total
    a cobrar (±1 €), mejor si está en «Cobro de facturas». Un mismo ingreso no vale para dos facturas."""
    hoy = hoy or date.today()
    pendientes = sorted((x for x in facturas if x.fecha_cobro is None), key=lambda x: (x.fecha, x.id))
    sugerencias = []
    if pendientes:
        todos = prevision.movimientos_tuyos(s, pendientes[0].fecha, hoy + timedelta(days=1))
        traspasos = prevision.ids_traspaso(todos)
        ingresos = sorted((m for m in todos if m.importe > 0 and m.id not in traspasos and m.tipo != "transferencia"),
                          key=lambda m: (m.categoria != "Cobro de facturas", m.fecha, m.id))
        usados: set[int] = set()
        for x in pendientes:
            total = float(x.total_a_cobrar)
            m = next((m for m in ingresos if m.id not in usados and m.fecha >= x.fecha
                      and abs(m.importe - total) <= TOLERANCIA_COBRO), None)
            if m:
                usados.add(m.id)
                sugerencias.append({"factura_id": x.id, "movimiento_id": m.id, "fecha": f(m.fecha),
                                    "importe": round(m.importe, 2)})
    return {
        "por_cobrar": {"total": n(sum((x.total_a_cobrar for x in pendientes), CERO)), "facturas": len(pendientes),
                       "mas_antigua_dias": max((hoy - pendientes[0].fecha).days, 0) if pendientes else None},
        "sugerencias_cobro": sugerencias,
    }


def dias_pendiente(x: Factura, hoy: date | None = None) -> int | None:
    """Días desde la fecha de la factura si aún no está cobrada."""
    return None if x.fecha_cobro else max(((hoy or date.today()) - x.fecha).days, 0)


class CobroIn(BaseModel):
    fecha: date | None = None  # vacío: hoy


@router.post("/autonomo/facturas/{factura_id}/cobrada")
def marcar_cobrada(factura_id: int, datos: CobroIn, s: Session = SesionDB):
    x = _obtener(s, Factura, factura_id)
    x.fecha_cobro = datos.fecha or date.today()
    s.commit()
    return {"ok": True, "fecha_cobro": f(x.fecha_cobro)}


@router.post("/autonomo/cobros/conciliar")
def conciliar_cobros(s: Session = SesionDB):
    """Da por cobradas, en la fecha del ingreso, todas las facturas con un ingreso del banco que cuadra."""
    facturas = s.scalars(select(Factura)).all()
    sugerencias = estado_cobros(s, facturas)["sugerencias_cobro"]
    por_id = {x.id: x for x in facturas}
    for sug in sugerencias:
        por_id[sug["factura_id"]].fecha_cobro = date.fromisoformat(sug["fecha"])
    s.commit()
    return {"ok": True, "marcadas": len(sugerencias)}


def avisos(s: Session) -> list[dict]:
    """Facturas que llevan más de 45 días sin cobrar (para Inicio)."""
    hoy = date.today()
    tarde = s.scalars(select(Factura).where(Factura.fecha_cobro.is_(None),
                                            Factura.fecha < hoy - timedelta(days=DIAS_AVISO_COBRO))
                      .order_by(Factura.fecha, Factura.id)).all()
    if not tarde:
        return []
    if len(tarde) <= 2:
        return [{"nivel": "aviso", "texto": f"La factura {x.numero} ({x.cliente.nombre}) lleva {(hoy - x.fecha).days} días sin cobrar.",
                 "ir": "/ingresos"} for x in tarde]
    x = tarde[0]
    return [{"nivel": "aviso", "texto": f"{len(tarde)} facturas llevan más de {DIAS_AVISO_COBRO} días sin cobrar; "
                                        f"la más antigua, la {x.numero}, {(hoy - x.fecha).days} días.", "ir": "/ingresos"}]


# --- Numeración ---------------------------------------------------------------

def avisos_numeracion(facturas: list[Factura], maximo: int = 4) -> list[str]:
    """Las facturas por fecha, comprobando que cada número sigue al anterior según `facturacion.siguiente_numero`:
    así salen los saltos y las fechas fuera de orden respecto al número."""
    orden = sorted(facturas, key=lambda x: (x.fecha, x.id))
    avisos = []
    for a, b in zip(orden, orden[1:]):
        # Sin un contador aparte del año (p. ej. «2026» a secas) no hay nada que comprobar
        if not any(g != str(a.fecha.year) for g in re.findall(r"\d+", a.numero)):
            continue
        esperado = facturacion.siguiente_numero(a.numero, a.fecha.year, b.fecha.year)
        if b.numero.strip() != esperado:
            avisos.append(f"Tras la {a.numero} ({a.fecha:%d/%m}) tocaría la {esperado} y va la {b.numero} ({b.fecha:%d/%m}).")
    if len(avisos) > maximo:
        avisos = avisos[:maximo] + [f"Y {len(avisos) - maximo} saltos más en la numeración de este año."]
    return avisos


# --- Libro de facturas para la gestoría ----------------------------------------

def _hoja_por_trimestres(hoja, cabecera: list[str], filas: list[tuple[int, list]], anchos: list[int], anio: int) -> None:
    """Rellena una hoja con las filas (trimestre, valores) y una fila de totales por trimestre y otra del año;
    los totales suman las columnas numéricas (float)."""
    from openpyxl.styles import Font
    negrita = Font(bold=True)
    hoja.append(cabecera)
    for celda in hoja[1]:
        celda.font = negrita

    def total(texto: str, grupo: list[list]) -> None:
        fila = [texto] + [round(sum(v[i] for v in grupo if isinstance(v[i], float)), 2)
                          if any(isinstance(v[i], float) for v in grupo) else "" for i in range(1, len(cabecera))]
        hoja.append(fila)
        for celda in hoja[hoja.max_row]:
            celda.font = negrita

    for t in range(1, 5):
        grupo = [v for tri, v in filas if tri == t]
        if not grupo:
            continue
        for v in grupo:
            hoja.append(v)
        total(f"Total {t}T", grupo)
    total(f"Total {anio}", [v for _, v in filas])
    for col, ancho in zip("ABCDEFGHIJKLMN", anchos):
        hoja.column_dimensions[col].width = ancho
    hoja.freeze_panes = "A2"


@router.get("/autonomo/libro.xlsx")
def libro_facturas(anio: int, s: Session = SesionDB):
    """Libro de facturas emitidas y de gastos del año, en Excel, con totales por trimestre: lo que pide la gestoría."""
    from openpyxl import Workbook
    facturas, gastos = _del_anio(s, Factura, anio), _del_anio(s, GastoAutonomo, anio)
    libro = Workbook()
    emitidas = libro.active
    emitidas.title = "Emitidas"
    _hoja_por_trimestres(
        emitidas, ["Número", "Fecha", "Cliente", "NIF", "Concepto", "Base", "% IVA", "Cuota IVA", "% retención",
                   "Retención", "Total", "Fecha de cobro"],
        [(autonomo.trimestre_de(x.fecha), [x.numero, x.fecha, x.cliente.nombre, x.cliente.nif or "", x.concepto, float(x.base),
                                           float(x.tipo_iva), float(x.cuota_iva), float(x.tipo_retencion),
                                           float(x.retencion), float(x.total_a_cobrar), x.fecha_cobro])
         for x in facturas],
        [12, 12, 28, 14, 40, 12, 8, 12, 11, 12, 12, 14], anio)
    hoja = libro.create_sheet("Gastos")
    _hoja_por_trimestres(
        hoja, ["Fecha", "Proveedor", "Concepto", "Categoría", "Base", "% IVA", "Cuota IVA", "% deducible",
               "Base deducible", "IVA deducible"],
        [(autonomo.trimestre_de(g.fecha), [g.fecha, g.proveedor, g.concepto, g.categoria, float(g.base), float(g.tipo_iva),
                                           float(g.cuota_iva), float(g.deducible_pct),
                                           float((g.base * g.deducible_pct / 100).quantize(CENT)),
                                           float((g.cuota_iva * g.deducible_pct / 100).quantize(CENT))])
         for g in gastos],
        [12, 28, 40, 14, 12, 8, 12, 11, 14, 14], anio)
    salida = io.BytesIO()
    libro.save(salida)
    return Response(salida.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="libro-facturas-{anio}.xlsx"'})


# --- Resumen del año: lo que hace falta para el 390 y el 347 --------------------

TIPOS_IVA = (21.0, 10.0, 4.0, 0.0)


@router.get("/autonomo/anual")
def resumen_anual(anio: int, s: Session = SesionDB):
    facturas, gastos = _del_anio(s, Factura, anio), _del_anio(s, GastoAutonomo, anio)
    por_tipo: dict[float, list[Decimal]] = {t: [CERO, CERO] for t in TIPOS_IVA}
    for x in facturas:
        fila = por_tipo.setdefault(float(x.tipo_iva), [CERO, CERO])
        fila[0] += x.base
        fila[1] += x.cuota_iva
    trimestres = []
    for t in range(1, 5):
        m = autonomo.calcular_303(anio, t, facturas, gastos)
        trimestres.append({"trimestre": t, "calculado": n(m.resultado),
                           "presentado": declaraciones.importe_presentado(s, "303", anio, f"{t}T")})
    presentados = [t["presentado"] for t in trimestres if t["presentado"] is not None]
    # 347: clientes y proveedores con más de 3.005,06 € en el año, IVA incluido. Las operaciones con retención
    # las declara tu cliente en el 190, no van en tu 347: se marcan para que lo sepas
    clientes: dict[int, dict] = {}
    for x in facturas:
        c = clientes.setdefault(x.cliente_id, {"nombre": x.cliente.nombre, "nif": x.cliente.nif or "", "tipo": "cliente",
                                               "importe": CERO, "operaciones": 0, "con_retencion": False})
        c["importe"] += x.base + x.cuota_iva
        c["operaciones"] += 1
        c["con_retencion"] = c["con_retencion"] or x.tipo_retencion > 0
    proveedores: dict[str, dict] = {}
    for g in gastos:
        nombre = (g.proveedor or g.concepto or "Sin proveedor").strip()
        p = proveedores.setdefault(nombre.lower(), {"nombre": nombre, "nif": "", "tipo": "proveedor", "importe": CERO,
                                                    "operaciones": 0, "con_retencion": False})
        p["importe"] += g.base + g.cuota_iva
        p["operaciones"] += 1
    terceros = [{**t, "importe": n(t["importe"])} for t in list(clientes.values()) + list(proveedores.values())
                if t["importe"] > UMBRAL_347]
    return {
        "anio": anio, "facturas": len(facturas), "gastos": len(gastos),
        "repercutido": [{"tipo": t, "base": n(b), "cuota": n(c)} for t, (b, c) in sorted(por_tipo.items(), reverse=True)],
        "base_total": n(sum((x.base for x in facturas), CERO)),
        "iva_repercutido": n(sum((x.cuota_iva for x in facturas), CERO)),
        "base_soportada": n(sum(((g.base * g.deducible_pct / 100).quantize(CENT) for g in gastos), CERO)),
        "iva_soportado": n(sum(((g.cuota_iva * g.deducible_pct / 100).quantize(CENT) for g in gastos if g.tipo_iva > 0), CERO)),
        "retenciones": n(sum((x.retencion for x in facturas), CERO)),
        "trimestres": trimestres,
        "presentado_303": round(sum(presentados), 2) if presentados else None,
        "calculado_303": round(sum(t["calculado"] for t in trimestres), 2),
        "umbral_347": float(UMBRAL_347),
        "terceros_347": sorted(terceros, key=lambda t: -t["importe"]),
    }


# --- Cuota de autónomos que carga el banco, como gasto ---------------------------

def cargos_tgss_sin_apuntar(s: Session, anio: int) -> list[dict]:
    """Cargos de la Seguridad Social en el banco ese año que no tienen ya un gasto de la categoría «cuota de
    autónomos» con la misma fecha (±3 días) y el mismo importe."""
    cargos = prevision.cargos_cuota_autonomos(s, date(anio, 1, 1), date(anio, 12, 31))
    margen = timedelta(days=DIAS_CUOTA)
    libres = s.scalars(select(GastoAutonomo).where(GastoAutonomo.categoria == "cuota_reta",
                                                   GastoAutonomo.fecha >= date(anio, 1, 1) - margen,
                                                   GastoAutonomo.fecha <= date(anio, 12, 31) + margen)).all()
    sin_apuntar = []
    for fecha, importe in sorted(cargos):
        g = next((g for g in libres if abs((g.fecha - fecha).days) <= DIAS_CUOTA and abs(float(g.base) - importe) < 0.01), None)
        if g is not None:
            libres.remove(g)
        else:
            sin_apuntar.append({"fecha": f(fecha), "importe": round(importe, 2)})
    return sin_apuntar


def _resumen_cargos(cargos: list[dict]) -> dict:
    return {"cargos": cargos, "n": len(cargos), "total": round(sum(c["importe"] for c in cargos), 2)}


@router.get("/autonomo/cuota-tgss")
def ver_cuota_tgss(anio: int, s: Session = SesionDB):
    return _resumen_cargos(cargos_tgss_sin_apuntar(s, anio))


@router.post("/autonomo/cuota-tgss")
def apuntar_cuota_tgss(anio: int, s: Session = SesionDB):
    """Apunta cada cargo como gasto de la actividad: cuota de autónomos, sin IVA y deducible al 100 %."""
    cargos = cargos_tgss_sin_apuntar(s, anio)
    for c in cargos:
        fecha = date.fromisoformat(c["fecha"])
        s.add(GastoAutonomo(fecha=fecha, proveedor="Seguridad Social (TGSS)",
                            concepto=f"Cuota de autónomos de {MESES[fecha.month - 1]}", categoria="cuota_reta",
                            base=Decimal(str(c["importe"])), tipo_iva=CERO, deducible_pct=Decimal("100")))
    s.commit()
    return {"ok": True, **_resumen_cargos(cargos)}


# --- Editar un gasto --------------------------------------------------------------

@router.put("/autonomo/gastos/{gasto_id}")
def actualizar_gasto_autonomo(gasto_id: int, datos: GastoAutonomoIn, s: Session = SesionDB):
    g = _obtener(s, GastoAutonomo, gasto_id)
    for campo, valor in datos.model_dump().items():
        setattr(g, campo, valor)
    s.commit()
    return {"ok": True}
