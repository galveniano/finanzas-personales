"""Importador de extractos descargados de la web de Banco Sabadell (Excel o CSV).

El formato exacto del extracto puede variar, así que se busca la fila de cabecera
por nombre de columna (fecha operativa, concepto, fecha valor, importe, saldo).
"""
import csv
import hashlib
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from finanzas.categorizar import categorizar
from finanzas.models import Cuenta, Movimiento

COLUMNAS = {
    "fecha": ("f. operativa", "fecha operativa", "fecha operación", "fecha operacion", "fecha"),
    "concepto": ("concepto", "descripción", "descripcion"),
    "fecha_valor": ("f. valor", "fecha valor"),
    "importe": ("importe", "importe (€)", "importe eur"),
    "saldo": ("saldo", "saldo (€)"),
}


MILES_CON_PUNTO = re.compile(r"^-?\d{1,3}(\.\d{3})+$")


@dataclass
class FilaExtracto:
    fecha: date
    concepto: str
    importe: Decimal
    fecha_valor: date | None = None
    saldo: Decimal | None = None


@dataclass
class ResultadoImportacion:
    leidos: int = 0
    nuevos: int = 0
    duplicados: int = 0


def parse_importe(valor) -> Decimal | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return Decimal(str(valor)).quantize(Decimal("0.01"))
    s = str(valor).strip().replace("€", "").replace("EUR", "").replace(" ", "").replace("\xa0", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif MILES_CON_PUNTO.match(s):  # «1.000» o «-1.500» son miles; «1234.56» o «1.5» siguen siendo decimales
        s = s.replace(".", "")
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def parse_fecha(valor) -> date | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = str(valor).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _norm(celda) -> str:
    return str(celda or "").strip().lower()


def _mapear_cabecera(fila: list) -> dict[str, int] | None:
    nombres = [_norm(c) for c in fila]
    mapa = {}
    for campo, alias in COLUMNAS.items():
        for i, n in enumerate(nombres):
            if n in alias and i not in mapa.values():
                mapa[campo] = i
                break
    if {"fecha", "concepto", "importe"} <= mapa.keys():
        return mapa
    return None


def _leer_filas(nombre: str, contenido: bytes) -> list[list]:
    ext = Path(nombre).suffix.lower()
    if ext == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    if ext == ".xls":
        import xlrd

        libro = xlrd.open_workbook(file_contents=contenido)
        hoja = libro.sheet_by_index(0)
        filas = []
        for r in range(hoja.nrows):
            fila = []
            for c in range(hoja.ncols):
                celda = hoja.cell(r, c)
                if celda.ctype == xlrd.XL_CELL_DATE:
                    fila.append(xlrd.xldate_as_datetime(celda.value, libro.datemode))
                else:
                    fila.append(celda.value)
            filas.append(fila)
        return filas
    # CSV: Sabadell suele usar ';' y latin-1
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    # El separador más frecuente gana (las líneas de cabecera del banco confunden a csv.Sniffer)
    separador = max(";,\t", key=texto.count)
    return [fila for fila in csv.reader(io.StringIO(texto), delimiter=separador)]


def parsear(nombre: str, contenido: bytes) -> list[FilaExtracto]:
    filas = _leer_filas(nombre, contenido)
    mapa = None
    resultado = []
    for fila in filas:
        if mapa is None:
            mapa = _mapear_cabecera(fila)
            continue
        if len(fila) <= max(mapa.values()):
            continue
        fecha = parse_fecha(fila[mapa["fecha"]])
        importe = parse_importe(fila[mapa["importe"]])
        if fecha is None or importe is None:
            continue
        resultado.append(FilaExtracto(
            fecha=fecha,
            concepto=str(fila[mapa["concepto"]] or "").strip(),
            importe=importe,
            fecha_valor=parse_fecha(fila[mapa["fecha_valor"]]) if "fecha_valor" in mapa else None,
            saldo=parse_importe(fila[mapa["saldo"]]) if "saldo" in mapa else None,
        ))
    if mapa is None:
        raise ValueError(
            "No encuentro la cabecera del extracto (busco columnas como Fecha, Concepto e Importe)."
        )
    return resultado


def huella(f: FilaExtracto, ocurrencia: int) -> str:
    base = f"{f.fecha}|{f.concepto}|{f.importe}|{f.saldo}|{ocurrencia}"
    return hashlib.sha256(base.encode()).hexdigest()[:32]


def importar(session: Session, cuenta: Cuenta, nombre: str, contenido: bytes) -> ResultadoImportacion:
    filas = parsear(nombre, contenido)
    res = ResultadoImportacion(leidos=len(filas))
    existentes = set(session.scalars(select(Movimiento.huella).where(Movimiento.cuenta_id == cuenta.id)))
    vistos: dict[tuple, int] = {}
    for f in filas:
        clave = (f.fecha, f.concepto, f.importe, f.saldo)
        vistos[clave] = vistos.get(clave, 0) + 1
        h = huella(f, vistos[clave])
        if h in existentes:
            res.duplicados += 1
            continue
        mov = Movimiento(cuenta_id=cuenta.id, fecha=f.fecha, fecha_valor=f.fecha_valor,
                         concepto=f.concepto, importe=f.importe, saldo=f.saldo, huella=h)
        mov.categoria_id = categorizar(session, f.concepto)
        session.add(mov)
        existentes.add(h)
        res.nuevos += 1
    # El saldo de la cuenta pasa a ser el del movimiento más reciente con saldo
    con_saldo = [f for f in filas if f.saldo is not None]
    if con_saldo:
        # Con varios movimientos el mismo día, el último depende del orden del extracto
        # (Sabadell suele listar del más reciente al más antiguo).
        descendente = con_saldo[0].fecha > con_saldo[-1].fecha
        fecha_max = max(f.fecha for f in con_saldo)
        del_dia = [f for f in con_saldo if f.fecha == fecha_max]
        ultimo = del_dia[0] if descendente else del_dia[-1]
        if cuenta.saldo_fecha is None or ultimo.fecha >= cuenta.saldo_fecha:
            cuenta.saldo, cuenta.saldo_fecha = ultimo.saldo, ultimo.fecha
    if cuenta.origen == "manual":
        cuenta.origen = "csv"
    session.commit()
    return res
